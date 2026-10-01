from collections import deque
from datetime import datetime, timedelta, date
import calendar
import json

from dateutil.relativedelta import relativedelta

from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils.timezone import now
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from docx import Document
from docx.shared import Pt

from .models import EventDutyRecord, Event, User, NFDQueue
from .services.event_service import get_events_for_date


# ============================================================
# Helper Functions
# ============================================================

def is_weekend(duty_date):
    """Return True if the given date is Saturday or Sunday."""
    return duty_date.weekday() >= 5


def get_previous_month(year, month):
    if month == 1:
        return year - 1, 12

    return year, month - 1


def get_month_dates(year, month):
    number_of_days = calendar.monthrange(year, month)[1]

    return [
        date(year, month, day)
        for day in range(1, number_of_days + 1)
    ]


def get_weekend_blocked_employees(year, month):
    """
    Return employee IDs who performed weekend duty
    in the previous two months.
    """
    blocked_employee_ids = set()

    current_year = year
    current_month = month

    for _ in range(2):
        current_year, current_month = get_previous_month(
            current_year,
            current_month,
        )

        previous_duties = EventDutyRecord.objects.filter(
            duty_date__year=current_year,
            duty_date__month=current_month,
        )

        for duty in previous_duties:
            if is_weekend(duty.duty_date):
                blocked_employee_ids.add(duty.employee_id)

    return blocked_employee_ids


def get_event_for_date(duty_date):
    """
    Return the first active event associated with the given date.
    """
    return get_events_for_date(duty_date).first()


def create_duty(employee, duty_date):
    """
    Create a duty and automatically attach the event
    associated with that date.
    """
    event = get_event_for_date(duty_date)

    return EventDutyRecord.objects.create(
        employee=employee,
        duty_date=duty_date,
        event=event,
    )


def parse_duty_date(value):
    """Safely convert YYYY-MM-DD text into a date object."""
    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d",
        ).date()
    except (TypeError, ValueError):
        return None


def get_active_employee(employee_id):
    """Return an active employee or None."""
    return User.objects.filter(
        id=employee_id,
        role="employee",
        status="active",
    ).first()


# ============================================================
# Superintendent Dashboard
# ============================================================

def suprident_dashboard(request):
    if request.session.get("role") != "suprident":
        messages.error(
            request,
            "You are not authorized to view this page.",
        )
        return redirect("login")

    current_year = now().year

    duties_qs = (
        EventDutyRecord.objects
        .select_related("employee", "event")
        .filter(
            duty_date__year=current_year,
        )
        .order_by("duty_date")
    )

    events_qs = (
        Event.objects
        .filter(
            date__year=current_year,
            enabled=True,
        )
        .order_by("date")
    )

    calendar_events = [
        {
            "title": duty.employee.name,
            "start": duty.duty_date.strftime("%Y-%m-%d"),
            "color": "#4F46E5",
            "allDay": True,
            "extendedProps": {
                "hr_number": duty.employee.hr_number,
                "designation": duty.employee.designation,
                "event": (
                    duty.event.name
                    if duty.event
                    else "Duty"
                ),
            },
        }
        for duty in duties_qs
    ]

    calendar_events += [
        {
            "title": event.name,
            "start": event.date.strftime("%Y-%m-%d"),
            "color": "#16A34A",
            "allDay": True,
            "extendedProps": {
                "event": "Event",
            },
        }
        for event in events_qs
    ]

    duties_table = [
        {
            "employee_name": duty.employee.name,
            "duty_date": duty.duty_date.strftime("%Y-%m-%d"),
            "event_name": (
                duty.event.name
                if duty.event
                else "—"
            ),
        }
        for duty in duties_qs
    ]

    return render(
        request,
        "core/super/dashboard.html",
        {
            "duties": duties_table,
            "calendar_events": calendar_events,
            "year": current_year,
        },
    )


# ============================================================
# Assign Duty
# ============================================================

def assign_duty(request):
    if request.session.get("role") != "suprident":
        messages.error(
            request,
            "Unauthorized access",
        )
        return redirect("login")

    employees = (
        User.objects
        .filter(
            role="employee",
            status="active",
        )
        .order_by("id")
    )

    selected_month_str = request.GET.get("month")

    if selected_month_str:
        try:
            selected_month = datetime.strptime(
                selected_month_str,
                "%Y-%m",
            ).date()

            selected_month = selected_month.replace(
                day=1
            )

        except ValueError:
            selected_month = now().date().replace(
                day=1
            )
    else:
        selected_month = now().date().replace(
            day=1
        )

    year = selected_month.year
    month = selected_month.month

    selected_month_label = selected_month.strftime(
        "%B %Y"
    )

    next_month = (
        selected_month.replace(day=28)
        + timedelta(days=4)
    ).replace(day=1)

    # --------------------------------------------------------
    # NFD EMPLOYEES FOR SELECTED MONTH
    #
    # IMPORTANT:
    # NFD records are NOT converted into duties here.
    # They remain in NFDQueue so they can appear in reports.
    # --------------------------------------------------------

    nfd_employee_ids = set(
        NFDQueue.objects
        .filter(
            assigned_month=selected_month,
        )
        .values_list(
            "employee_id",
            flat=True,
        )
    )

    # --------------------------------------------------------
    # Get Duties for Selected Month
    # --------------------------------------------------------

    duties = (
        EventDutyRecord.objects
        .filter(
            duty_date__year=year,
            duty_date__month=month,
        )
        .select_related(
            "employee",
            "event",
        )
        .order_by("duty_date")
    )

    # --------------------------------------------------------
    # Auto Assign
    # --------------------------------------------------------

    auto_assign_flag = (
        request.GET.get("auto_assign") == "1"
    )

    if auto_assign_flag:

        all_dates = get_month_dates(
            year,
            month,
        )

        assigned_dates = {
            duty.duty_date
            for duty in duties
        }

        # ----------------------------------------------------
        # Exclude employees already present in NFD
        # ----------------------------------------------------

        available_employees = (
            User.objects
            .filter(
                role="employee",
                status="active",
            )
            .exclude(
                id__in=nfd_employee_ids,
            )
            .order_by("id")
        )

        print(
            "Available Employees:",
            [
                employee.name
                for employee in available_employees
            ],
        )

        # ----------------------------------------------------
        # Weekend restriction
        # ----------------------------------------------------

        two_months_ago = (
            selected_month
            - timedelta(days=60)
        ).replace(day=1)

        weekend_blocked_ids = set()

        for employee in available_employees:

            recent_duties = (
                EventDutyRecord.objects.filter(
                    employee=employee,
                    duty_date__gte=two_months_ago,
                    duty_date__lt=selected_month,
                )
            )

            had_weekend = any(
                is_weekend(duty.duty_date)
                for duty in recent_duties
            )

            if had_weekend:
                weekend_blocked_ids.add(
                    employee.id
                )

        print(
            "Weekend Blocked:",
            weekend_blocked_ids,
        )

        employee_queue = deque(
            available_employees
        )

        # ----------------------------------------------------
        # Assign one employee per date
        # ----------------------------------------------------

        for duty_date in all_dates:

            if duty_date in assigned_dates:
                continue

            weekend = is_weekend(duty_date)
            assigned = False
            attempts = 0

            while (
                employee_queue
                and attempts < len(employee_queue)
            ):

                employee = employee_queue.popleft()
                employee_queue.append(employee)

                attempts += 1

                # --------------------------------------------
                # One duty per employee per month
                # --------------------------------------------

                if employee.id in {
                    duty.employee_id
                    for duty in duties
                }:
                    continue

                # --------------------------------------------
                # Weekend restriction
                # --------------------------------------------

                if (
                    weekend
                    and employee.id in weekend_blocked_ids
                ):
                    continue

                create_duty(
                    employee,
                    duty_date,
                )

                print(
                    f"Auto Assigned: "
                    f"{employee.name} -> "
                    f"{duty_date}"
                )

                assigned_dates.add(
                    duty_date
                )

                assigned = True
                break

            if not assigned:
                print(
                    f"No available employee "
                    f"for {duty_date}"
                )

        # ----------------------------------------------------
        # Refresh duties
        # ----------------------------------------------------

        duties = (
            EventDutyRecord.objects
            .filter(
                duty_date__year=year,
                duty_date__month=month,
            )
            .select_related(
                "employee",
                "event",
            )
            .order_by("duty_date")
        )

    # --------------------------------------------------------
    # Month Selector
    # --------------------------------------------------------

    start_year = now().year - 1
    end_year = now().year + 5

    month_options = [
        (
            datetime(y, m, 1).strftime("%Y-%m"),
            datetime(y, m, 1).strftime("%B %Y"),
        )
        for y in range(
            start_year,
            end_year + 1,
        )
        for m in range(1, 13)
    ]

    # --------------------------------------------------------
    # NFDs for NEXT MONTH
    #
    # These are displayed on the Assign Duty page.
    # --------------------------------------------------------

    nfds = (
        NFDQueue.objects
        .filter(
            assigned_month=next_month,
        )
        .select_related("employee")
        .order_by(
            "duty_date",
            "employee__name",
        )
    )

    return render(
        request,
        "core/super/assign_duty.html",
        {
            "employees": employees,
            "duties": duties,
            "selected_month": selected_month.strftime(
                "%Y-%m"
            ),
            "selected_month_label": selected_month_label,
            "month_options": month_options,
            "nfds": nfds,
        },
    )


# ============================================================
# Auto Assign Monthly Duties
# ============================================================

def auto_assign_monthly_duties(request):
    if request.session.get("role") != "suprident":
        messages.error(
            request,
            "Unauthorized access",
        )
        return redirect("login")

    if request.method != "POST":
        messages.error(
            request,
            "Invalid request method.",
        )
        return redirect("assign_duty")

    selected_month = request.POST.get("month")

    if not selected_month:
        messages.error(
            request,
            "Please select a valid month.",
        )
        return redirect("assign_duty")

    try:
        year, month = map(
            int,
            selected_month.split("-"),
        )
    except ValueError:
        messages.error(
            request,
            "Invalid month format.",
        )
        return redirect("assign_duty")

    month_dates = get_month_dates(
        year,
        month,
    )

    prev_year, prev_month = get_previous_month(
        year,
        month,
    )

    all_employees = list(
        User.objects
        .filter(
            role="employee",
            status="active",
        )
        .order_by("id")
    )

    prev_duties = EventDutyRecord.objects.filter(
        duty_date__year=prev_year,
        duty_date__month=prev_month,
    )

    assigned_last_month = {
        duty.employee_id
        for duty in prev_duties
    }

    all_ids = {
        employee.id
        for employee in all_employees
    }

    nfd_employee_ids = list(
        all_ids - assigned_last_month
    )

    priority_employees = (
        list(
            User.objects.filter(
                id__in=nfd_employee_ids
            ).order_by("id")
        )
        +
        list(
            User.objects.filter(
                id__in=assigned_last_month
            ).order_by("id")
        )
    )

    # --------------------------------------------------------
    # Delete existing duties for this month
    # --------------------------------------------------------

    EventDutyRecord.objects.filter(
        duty_date__year=year,
        duty_date__month=month,
    ).delete()

    blocked_weekend_ids = (
        get_weekend_blocked_employees(
            year,
            month,
        )
    )

    assigned_employee_ids = set()
    duties_to_create = []

    # --------------------------------------------------------
    # Generate Duties
    # --------------------------------------------------------

    for duty_date in month_dates:

        assigned = False

        for employee in priority_employees:

            # One duty per employee per month
            if employee.id in assigned_employee_ids:
                continue

            # Weekend restriction
            if (
                is_weekend(duty_date)
                and employee.id in blocked_weekend_ids
            ):
                continue

            event = get_event_for_date(
                duty_date
            )

            duties_to_create.append(
                EventDutyRecord(
                    employee=employee,
                    duty_date=duty_date,
                    event=event,
                )
            )

            assigned_employee_ids.add(
                employee.id
            )

            assigned = True
            break

        # ----------------------------------------------------
        # Fallback
        #
        # IMPORTANT:
        # Still respect:
        # 1. One duty per employee/month
        # 2. Weekend restriction
        # ----------------------------------------------------

        if not assigned:

            for employee in all_employees:

                if employee.id in assigned_employee_ids:
                    continue

                if (
                    is_weekend(duty_date)
                    and employee.id in blocked_weekend_ids
                ):
                    continue

                event = get_event_for_date(
                    duty_date
                )

                duties_to_create.append(
                    EventDutyRecord(
                        employee=employee,
                        duty_date=duty_date,
                        event=event,
                    )
                )

                assigned_employee_ids.add(
                    employee.id
                )

                assigned = True
                break

    # --------------------------------------------------------
    # Save generated duties
    # --------------------------------------------------------

    EventDutyRecord.objects.bulk_create(
        duties_to_create
    )

    # --------------------------------------------------------
    # Calculate NFD employees
    # --------------------------------------------------------

    all_assigned = {
        duty.employee_id
        for duty in duties_to_create
    }

    new_nfd_ids = list(
        all_ids - all_assigned
    )

    selected_month_date = date(
        year,
        month,
        1,
    )

    # --------------------------------------------------------
    # IMPORTANT:
    #
    # Remove old NFD records for this SAME month only,
    # then regenerate the correct NFD list.
    #
    # NFD records are NOT deleted later by assign_duty()
    # or generate_nfd_for_month().
    # --------------------------------------------------------

    NFDQueue.objects.filter(
        assigned_month=selected_month_date
    ).delete()

    available_nfd_dates = list(
        month_dates
    )

    # --------------------------------------------------------
    # Create NFD records
    # --------------------------------------------------------

    for employee_id in new_nfd_ids:

        if available_nfd_dates:

            duty_date = available_nfd_dates.pop(0)

            NFDQueue.objects.create(
                employee_id=employee_id,
                assigned_id=(
                    f"NFD-{employee_id}-{duty_date}"
                ),
                assigned_month=selected_month_date,
                duty_type=(
                    "weekend"
                    if is_weekend(duty_date)
                    else "weekday"
                ),
                duty_date=duty_date,
                added_to_queue_at=now(),
            )

        else:

            NFDQueue.objects.create(
                employee_id=employee_id,
                assigned_id=(
                    f"NFD-{employee_id}-"
                    f"{selected_month_date}"
                ),
                assigned_month=selected_month_date,
                duty_type="weekday",
                duty_date=None,
                added_to_queue_at=now(),
            )

    messages.success(
        request,
        f"Duties for "
        f"{calendar.month_name[month]} "
        f"{year} auto-generated.",
    )

    return redirect(
        f"/assign-duty/?month={selected_month}"
    )


# ============================================================
# Generate NFD
# ============================================================

@csrf_protect
@require_POST
def generate_nfd_for_month(request):

    try:

        selected_month = request.POST.get("month")

        if not selected_month:
            return JsonResponse(
                {
                    "status": "error",
                    "message": "Month is required.",
                }
            )

        year, month = map(
            int,
            selected_month.split("-"),
        )

        start_of_month = date(
            year,
            month,
            1,
        )

        end_of_month = date(
            year,
            month,
            calendar.monthrange(
                year,
                month,
            )[1],
        )

        # ----------------------------------------------------
        # DO NOT CONVERT NFDs INTO DUTIES
        #
        # Existing NFDQueue records must remain available
        # for reporting.
        # ----------------------------------------------------

        existing_nfd_count = (
            NFDQueue.objects
            .filter(
                assigned_month=start_of_month,
            )
            .count()
        )

        # ----------------------------------------------------
        # Find employees who have no duty in selected month
        # ----------------------------------------------------

        all_employees = User.objects.filter(
            role="employee",
            status="active",
        )

        assigned_ids = set(
            EventDutyRecord.objects.filter(
                duty_date__range=(
                    start_of_month,
                    end_of_month,
                )
            ).values_list(
                "employee_id",
                flat=True,
            )
        )

        unassigned_employees = (
            all_employees.exclude(
                id__in=assigned_ids,
            )
        )

        # ----------------------------------------------------
        # Calculate next month
        # ----------------------------------------------------

        next_month = (
            start_of_month.replace(day=28)
            + timedelta(days=4)
        ).replace(day=1)

        next_year = next_month.year
        next_month_number = next_month.month

        days_in_next_month = calendar.monthrange(
            next_year,
            next_month_number,
        )[1]

        # ----------------------------------------------------
        # Existing next-month NFD records
        #
        # DO NOT delete them.
        # ----------------------------------------------------

        used_dates = set(
            NFDQueue.objects
            .filter(
                assigned_month=next_month,
            )
            .exclude(
                duty_date=None,
            )
            .values_list(
                "duty_date",
                flat=True,
            )
        )

        already_in_nfd = set(
            NFDQueue.objects
            .filter(
                assigned_month=next_month,
            )
            .values_list(
                "employee_id",
                flat=True,
            )
        )

        available_dates = [
            date(
                next_year,
                next_month_number,
                day,
            )
            for day in range(
                1,
                days_in_next_month + 1,
            )
            if date(
                next_year,
                next_month_number,
                day,
            ) not in used_dates
        ]

        added = 0

        # ----------------------------------------------------
        # Add new NFD employees for next month
        # ----------------------------------------------------

        for employee in unassigned_employees:

            if employee.id in already_in_nfd:
                continue

            if not available_dates:
                NFDQueue.objects.create(
                    employee=employee,
                    assigned_id=(
                        f"NFD-{employee.id}-"
                        f"{next_month}"
                    ),
                    assigned_month=next_month,
                    duty_type="weekday",
                    duty_date=None,
                    added_to_queue_at=now(),
                )

                added += 1
                continue

            duty_date = available_dates.pop(0)

            duty_type = (
                "weekend"
                if is_weekend(duty_date)
                else "weekday"
            )

            NFDQueue.objects.create(
                employee=employee,
                assigned_id=(
                    f"NFD-{employee.id}-{duty_date}"
                ),
                assigned_month=next_month,
                duty_type=duty_type,
                duty_date=duty_date,
                added_to_queue_at=now(),
            )

            added += 1

        return JsonResponse(
            {
                "status": "success",
                "message": (
                    f"{existing_nfd_count} existing NFDs "
                    f"for {start_of_month.strftime('%B %Y')} "
                    f"were preserved. "
                    f"{added} employees added to NFD "
                    f"for {next_month.strftime('%B %Y')}."
                ),
            }
        )

    except Exception as e:

        return JsonResponse(
            {
                "status": "error",
                "message": str(e),
            },
            status=500,
        )


# ============================================================
# Delete NFD
# ============================================================

@require_POST
def delete_nfd(request, nfd_id):

    if request.session.get("role") != "suprident":
        return JsonResponse(
            {
                "status": "error",
                "message": "Unauthorized",
            },
            status=403,
        )

    nfd = get_object_or_404(
        NFDQueue,
        id=nfd_id,
    )

    nfd.delete()

    messages.success(
        request,
        "NFD entry deleted successfully.",
    )

    return redirect(
        request.META.get(
            "HTTP_REFERER",
            "assign_duty",
        )
    )


# ============================================================
# Edit Duty
# ============================================================

@require_POST
def edit_duty_ajax(request, duty_id):

    if request.session.get("role") != "suprident":
        return JsonResponse(
            {
                "status": "error",
                "message": "Unauthorized",
            },
            status=403,
        )

    try:

        data = json.loads(request.body)

        date_str = data.get("date")
        employee_id = data.get("employee")

        if not date_str or not employee_id:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Date and employee "
                        "are required."
                    ),
                }
            )

        duty_date = parse_duty_date(
            date_str
        )

        if not duty_date:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Invalid date format. "
                        "Use YYYY-MM-DD."
                    ),
                }
            )

        employee = get_active_employee(
            employee_id
        )

        if not employee:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Selected employee "
                        "is not active or "
                        "does not exist."
                    ),
                }
            )

        duty = get_object_or_404(
            EventDutyRecord,
            id=duty_id,
        )

        # ----------------------------------------------------
        # Prevent another employee on same date
        # ----------------------------------------------------

        existing_duty = (
            EventDutyRecord.objects
            .filter(
                duty_date=duty_date,
            )
            .exclude(
                id=duty.id,
            )
            .first()
        )

        if existing_duty:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        f"{existing_duty.employee.name} "
                        f"is already assigned on "
                        f"{duty_date}. "
                        f"Only one employee per day "
                        f"is allowed."
                    ),
                }
            )

        # ----------------------------------------------------
        # One duty per employee per month
        # ----------------------------------------------------

        existing_monthly_duty = (
            EventDutyRecord.objects
            .filter(
                employee_id=employee.id,
                duty_date__year=duty_date.year,
                duty_date__month=duty_date.month,
            )
            .exclude(
                id=duty.id,
            )
            .first()
        )

        if existing_monthly_duty:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        f"This employee already "
                        f"has a duty in this month "
                        f"on "
                        f"{existing_monthly_duty.duty_date.strftime('%Y-%m-%d')}."
                    ),
                }
            )

        # ----------------------------------------------------
        # Weekend Rule
        # ----------------------------------------------------

        if is_weekend(duty_date):

            two_months_ago = (
                duty_date
                - relativedelta(months=2)
            )

            recent_weekend_duty = (
                EventDutyRecord.objects
                .filter(
                    employee_id=employee.id,
                    duty_date__gte=two_months_ago,
                    duty_date__lt=duty_date,
                )
                .filter(
                    duty_date__week_day__in=[1, 7],
                )
            )

            if recent_weekend_duty.exists():
                return JsonResponse(
                    {
                        "status": "error",
                        "message": (
                            "This employee has "
                            "already done a weekend "
                            "duty in the last 2 months. "
                            "Only weekdays are allowed "
                            "for now."
                        ),
                    }
                )

        # ----------------------------------------------------
        # Automatically Fetch Event
        # ----------------------------------------------------

        event = get_event_for_date(
            duty_date
        )

        duty.duty_date = duty_date
        duty.employee = employee
        duty.event = event

        duty.save()

        return JsonResponse(
            {
                "status": "success",
                "message": (
                    "Duty updated successfully."
                ),
                "event_name": (
                    event.name
                    if event
                    else "—"
                ),
            }
        )

    except json.JSONDecodeError:

        return JsonResponse(
            {
                "status": "error",
                "message": "Invalid JSON data.",
            }
        )

    except Exception as e:

        return JsonResponse(
            {
                "status": "error",
                "message": str(e),
            }
        )


# ============================================================
# Delete Duty
# ============================================================

@require_POST
def delete_duty_ajax(request, duty_id):

    if request.session.get("role") != "suprident":
        return JsonResponse(
            {
                "status": "error",
                "message": "Unauthorized",
            },
            status=403,
        )

    try:

        duty = get_object_or_404(
            EventDutyRecord,
            id=duty_id,
        )

        duty.delete()

        return JsonResponse(
            {
                "status": "success",
                "id": duty_id,
                "message": (
                    "Duty deleted successfully!"
                ),
            }
        )

    except Exception as e:

        return JsonResponse(
            {
                "status": "error",
                "message": str(e),
            },
            status=500,
        )


# ============================================================
# Save Duty
# ============================================================

@require_POST
def save_duty(request):

    if request.session.get("role") != "suprident":
        return JsonResponse(
            {
                "status": "error",
                "message": "Unauthorized",
            },
            status=403,
        )

    try:

        date_str = request.POST.get("date")
        employee_id = request.POST.get("employee")
        selected_month = request.POST.get(
            "selected_month"
        )

        if (
            not date_str
            or not employee_id
            or not selected_month
        ):
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Date, employee, and "
                        "selected month are required."
                    ),
                }
            )

        duty_date = parse_duty_date(
            date_str
        )

        if not duty_date:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Invalid date format. "
                        "Use YYYY-MM-DD."
                    ),
                }
            )

        # ----------------------------------------------------
        # Validate Selected Month
        # ----------------------------------------------------

        duty_month_str = duty_date.strftime(
            "%Y-%m"
        )

        if selected_month != duty_month_str:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Selected month and "
                        "duty date do not match."
                    ),
                }
            )

        # ----------------------------------------------------
        # Validate Employee
        # ----------------------------------------------------

        employee = get_active_employee(
            employee_id
        )

        if not employee:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "Selected employee "
                        "is not active or "
                        "does not exist."
                    ),
                }
            )

        # ----------------------------------------------------
        # Only One Employee Per Day
        # ----------------------------------------------------

        existing_duty = (
            EventDutyRecord.objects
            .filter(
                duty_date=duty_date,
            )
            .first()
        )

        if existing_duty:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        f"{existing_duty.employee.name} "
                        f"is already assigned on "
                        f"{duty_date}. "
                        f"Only one employee per day "
                        f"is allowed."
                    ),
                }
            )

        # ----------------------------------------------------
        # One Duty Per Employee Per Month
        # ----------------------------------------------------

        existing_monthly_duty = (
            EventDutyRecord.objects
            .filter(
                employee_id=employee.id,
                duty_date__year=duty_date.year,
                duty_date__month=duty_date.month,
            )
            .first()
        )

        if existing_monthly_duty:
            return JsonResponse(
                {
                    "status": "error",
                    "message": (
                        "This employee already "
                        "has a duty in this month "
                        f"on {existing_monthly_duty.duty_date.strftime('%Y-%m-%d')}."
                    ),
                }
            )

        # ----------------------------------------------------
        # Weekend Rule
        # ----------------------------------------------------

        if is_weekend(duty_date):

            two_months_ago = (
                duty_date
                - relativedelta(months=2)
            )

            recent_weekend_duty = (
                EventDutyRecord.objects
                .filter(
                    employee_id=employee.id,
                    duty_date__gte=two_months_ago,
                    duty_date__lt=duty_date,
                )
                .filter(
                    duty_date__week_day__in=[1, 7],
                )
            )

            if recent_weekend_duty.exists():
                return JsonResponse(
                    {
                        "status": "error",
                        "message": (
                            "This employee has "
                            "already done a weekend "
                            "duty in the last 2 months. "
                            "Only weekdays are allowed "
                            "for now."
                        ),
                    }
                )

        # ----------------------------------------------------
        # Automatically Fetch Event
        # ----------------------------------------------------

        event = get_event_for_date(
            duty_date
        )

        # ----------------------------------------------------
        # Create Duty
        # ----------------------------------------------------

        duty = EventDutyRecord.objects.create(
            duty_date=duty_date,
            employee=employee,
            event=event,
        )

        return JsonResponse(
            {
                "status": "success",
                "message": (
                    "Duty assigned successfully!"
                ),
                "id": duty.id,
                "duty_date": duty.duty_date.strftime(
                    "%Y-%m-%d"
                ),
                "employee_name": duty.employee.name,
                "event_name": (
                    event.name
                    if event
                    else "—"
                ),
            }
        )

    except Exception as e:

        return JsonResponse(
            {
                "status": "error",
                "message": str(e),
            },
            status=500,
        )


# ============================================================
# Generate Report Page
# ============================================================

def generate_report(request):
    """
    Display duty report and NFD list
    for the selected month.
    """

    if request.session.get("role") != "suprident":
        messages.error(
            request,
            "Unauthorized access",
        )
        return redirect("login")

    selected_month = request.GET.get(
        "month",
        "",
    ).strip()

    month_options = []
    current_year = now().year

    for month_number in range(1, 13):

        month_options.append(
            (
                f"{current_year}-{month_number:02d}",
                date(
                    current_year,
                    month_number,
                    1,
                ).strftime("%B %Y"),
            )
        )

    duties = EventDutyRecord.objects.none()
    nfds = NFDQueue.objects.none()
    selected_month_label = ""
    nfd_month_label = ""

    if selected_month:

        try:

            selected_year, selected_month_number = map(
                int,
                selected_month.split("-"),
            )

            selected_month_date = date(
                selected_year,
                selected_month_number,
                1,
            )

            selected_month_label = (
                selected_month_date.strftime(
                    "%B %Y"
                )
            )

            # ------------------------------------------------
            # Assigned Duties
            # ------------------------------------------------

            duties = (
                EventDutyRecord.objects
                .select_related(
                    "employee",
                    "event",
                )
                .filter(
                    duty_date__year=selected_year,
                    duty_date__month=selected_month_number,
                )
                .order_by(
                    "duty_date",
                    "employee__name",
                )
            )

            # ------------------------------------------------
            # NFD List
            #
            # The report for a month shows NFDs generated for
            # the NEXT month.
            #
            # Example:
            # September 2026 report -> October 2026 NFDs
            # October 2026 report   -> November 2026 NFDs
            #
            # NFDQueue.assigned_month identifies the month of
            # the NFD queue. duty_date is the actual NFD date.
            # ------------------------------------------------

            next_month_date = (
                selected_month_date.replace(day=28)
                + timedelta(days=4)
            ).replace(day=1)

            nfds = (
                NFDQueue.objects
                .select_related("employee")
                .filter(
                    assigned_month=next_month_date,
                )
                .order_by(
                    "duty_date",
                    "employee__name",
                )
            )

            nfd_month_label = next_month_date.strftime(
                "%B %Y"
            )

        except (ValueError, TypeError):

            selected_month = ""
            selected_month_label = ""
            nfds = NFDQueue.objects.none()

    context = {
        "duties": duties,
        "nfds": nfds,
        "month_options": month_options,
        "selected_month": selected_month,
        "selected_month_label": selected_month_label,
        "nfd_month_label": nfd_month_label,
        "can_download": bool(selected_month),
    }

    return render(
        request,
        "core/super/report.html",
        context,
    )


# ============================================================
# Download Report PDF
# ============================================================

def download_report_pdf(request):
    """
    Generate and download the selected month's
    duty report and NFD list as PDF.
    """

    if request.session.get("role") != "suprident":
        return HttpResponse(
            "Unauthorized access.",
            status=403,
        )

    selected_month = request.GET.get(
        "month",
        "",
    ).strip()

    if not selected_month:
        return HttpResponse(
            "Please select a month first.",
            status=400,
        )

    try:

        selected_year, selected_month_number = map(
            int,
            selected_month.split("-"),
        )

        selected_month_date = date(
            selected_year,
            selected_month_number,
            1,
        )

        selected_month_label = (
            selected_month_date.strftime(
                "%B %Y"
            )
        )

    except (ValueError, TypeError):

        return HttpResponse(
            "Invalid month.",
            status=400,
        )

    # --------------------------------------------------------
    # Assigned Duties
    # --------------------------------------------------------

    duties = list(
        EventDutyRecord.objects
        .select_related(
            "employee",
            "event",
        )
        .filter(
            duty_date__year=selected_year,
            duty_date__month=selected_month_number,
        )
        .order_by(
            "duty_date",
            "employee__name",
        )
    )

    # --------------------------------------------------------
    # NFD List
    # --------------------------------------------------------
    # The selected month's report displays the NEXT month's
    # NFD queue.
    #
    # Example:
    # September 2026 report -> October 2026 NFDs
    # October 2026 report   -> November 2026 NFDs
    # --------------------------------------------------------

    next_month_date = (
        selected_month_date.replace(day=28)
        + timedelta(days=4)
    ).replace(day=1)

    next_month_label = next_month_date.strftime(
        "%B %Y"
    )

    nfds = list(
        NFDQueue.objects
        .select_related("employee")
        .filter(
            assigned_month=next_month_date,
        )
        .order_by(
            "duty_date",
            "employee__name",
        )
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        f'attachment; '
        f'filename="Duty_Report_{selected_month}.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    elements = []

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            "Duty Roaster Report",
            styles["Title"],
        )
    )

    elements.append(
        Spacer(1, 10)
    )

    elements.append(
        Paragraph(
            f"Month: {selected_month_label}",
            styles["Heading2"],
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # ========================================================
    # DUTY DETAILS
    # ========================================================

    elements.append(
        Paragraph(
            "Duty Details",
            styles["Heading2"],
        )
    )

    elements.append(
        Spacer(1, 10)
    )

    table_data = [
        [
            "#",
            "Employee Name",
            "HR Number",
            "Designation",
            "Duty Date",
            "Event",
        ]
    ]

    for index, duty in enumerate(
        duties,
        start=1,
    ):

        table_data.append(
            [
                str(index),
                duty.employee.name,
                duty.employee.hr_number,
                duty.employee.designation,
                duty.duty_date.strftime(
                    "%d %b %Y"
                ),
                (
                    duty.event.name
                    if duty.event
                    else "—"
                ),
            ]
        )

    if not duties:

        table_data.append(
            [
                "",
                "No duties assigned",
                "",
                "",
                "",
                "",
            ]
        )

    table = Table(
        table_data,
        repeatRows=1,
        colWidths=[
            25,
            100,
            65,
            100,
            75,
            80,
        ],
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#e5e7eb"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    elements.append(table)

    # ========================================================
    # NFD LIST
    # ========================================================

    elements.append(
        Spacer(1, 20)
    )

    elements.append(
        Paragraph(
            f"NFD List - {next_month_label}",
            styles["Heading2"],
        )
    )

    elements.append(
        Spacer(1, 10)
    )

    nfd_table_data = [
        [
            "#",
            "Employee Name",
            "HR Number",
            "Designation",
            "NFD Date",
            "Duty Type",
        ]
    ]

    for index, nfd in enumerate(
        nfds,
        start=1,
    ):

        nfd_table_data.append(
            [
                str(index),
                nfd.employee.name,
                nfd.employee.hr_number,
                nfd.employee.designation,
                (
                    nfd.duty_date.strftime(
                        "%d %b %Y"
                    )
                    if nfd.duty_date
                    else "—"
                ),
                nfd.duty_type or "—",
            ]
        )

    if not nfds:

        nfd_table_data.append(
            [
                "",
                "No NFD employees",
                "",
                "",
                "",
                "",
            ]
        )

    nfd_table = Table(
        nfd_table_data,
        repeatRows=1,
        colWidths=[
            25,
            100,
            65,
            100,
            75,
            80,
        ],
    )

    nfd_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#fef3c7"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    elements.append(
        nfd_table
    )

    # --------------------------------------------------------
    # Generate PDF
    # --------------------------------------------------------

    document.build(elements)

    return response


# ============================================================
# Download Report Word
# ============================================================

def download_report_word(request):
    """
    Generate and download the selected month's
    duty report and NFD list as Word.
    """

    if request.session.get("role") != "suprident":
        return HttpResponse(
            "Unauthorized access.",
            status=403,
        )

    selected_month = request.GET.get(
        "month",
        "",
    ).strip()

    if not selected_month:
        return HttpResponse(
            "Please select a month first.",
            status=400,
        )

    try:

        selected_year, selected_month_number = map(
            int,
            selected_month.split("-"),
        )

        selected_month_date = date(
            selected_year,
            selected_month_number,
            1,
        )

        selected_month_label = (
            selected_month_date.strftime(
                "%B %Y"
            )
        )

    except (ValueError, TypeError):

        return HttpResponse(
            "Invalid month.",
            status=400,
        )

    # --------------------------------------------------------
    # Assigned Duties
    # --------------------------------------------------------

    duties = list(
        EventDutyRecord.objects
        .select_related(
            "employee",
            "event",
        )
        .filter(
            duty_date__year=selected_year,
            duty_date__month=selected_month_number,
        )
        .order_by(
            "duty_date",
            "employee__name",
        )
    )

    # --------------------------------------------------------
    # NFD List
    # --------------------------------------------------------
    # The selected month's report displays the NEXT month's
    # NFD queue.
    #
    # Example:
    # September 2026 report -> October 2026 NFDs
    # October 2026 report   -> November 2026 NFDs
    # --------------------------------------------------------

    next_month_date = (
        selected_month_date.replace(day=28)
        + timedelta(days=4)
    ).replace(day=1)

    next_month_label = next_month_date.strftime(
        "%B %Y"
    )

    nfds = list(
        NFDQueue.objects
        .select_related("employee")
        .filter(
            assigned_month=next_month_date,
        )
        .order_by(
            "duty_date",
            "employee__name",
        )
    )


    # ========================================================
    # Create Word Document
    # ========================================================

    document = Document()

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = document.add_heading(
        "Duty Roaster Report",
        level=1,
    )

    title.alignment = 1

    # --------------------------------------------------------
    # Month
    # --------------------------------------------------------

    month_paragraph = document.add_paragraph()

    month_run = month_paragraph.add_run(
        f"Month: {selected_month_label}"
    )

    month_run.bold = True

    document.add_paragraph()

    # ========================================================
    # DUTY DETAILS
    # ========================================================

    document.add_heading(
        "Duty Details",
        level=2,
    )

    table = document.add_table(
        rows=1,
        cols=6,
    )

    table.style = "Table Grid"

    headers = [
        "#",
        "Employee Name",
        "HR Number",
        "Designation",
        "Duty Date",
        "Event",
    ]

    header_cells = table.rows[0].cells

    for index, header in enumerate(headers):

        header_cells[index].text = header

        for paragraph in header_cells[index].paragraphs:

            for run in paragraph.runs:

                run.bold = True
                run.font.size = Pt(9)

    # --------------------------------------------------------
    # Duty Rows
    # --------------------------------------------------------

    for index, duty in enumerate(
        duties,
        start=1,
    ):

        row_cells = table.add_row().cells

        row_cells[0].text = str(index)
        row_cells[1].text = duty.employee.name
        row_cells[2].text = duty.employee.hr_number
        row_cells[3].text = duty.employee.designation
        row_cells[4].text = duty.duty_date.strftime(
            "%d %b %Y"
        )
        row_cells[5].text = (
            duty.event.name
            if duty.event
            else "—"
        )

    if not duties:

        row_cells = table.add_row().cells

        row_cells[0].text = ""
        row_cells[1].text = (
            "No duties assigned for this month."
        )
        row_cells[2].text = ""
        row_cells[3].text = ""
        row_cells[4].text = ""
        row_cells[5].text = ""

    # ========================================================
    # NFD LIST
    # ========================================================

    document.add_paragraph()

    document.add_heading(
        f"NFD List - {next_month_label}",
        level=2,
    )

    nfd_table = document.add_table(
        rows=1,
        cols=6,
    )

    nfd_table.style = "Table Grid"

    nfd_headers = [
        "#",
        "Employee Name",
        "HR Number",
        "Designation",
        "NFD Date",
        "Duty Type",
    ]

    nfd_header_cells = nfd_table.rows[0].cells

    for index, header in enumerate(nfd_headers):

        nfd_header_cells[index].text = header

        for paragraph in nfd_header_cells[index].paragraphs:

            for run in paragraph.runs:

                run.bold = True
                run.font.size = Pt(9)

    # --------------------------------------------------------
    # NFD Rows
    # --------------------------------------------------------

    for index, nfd in enumerate(
        nfds,
        start=1,
    ):

        row_cells = nfd_table.add_row().cells

        row_cells[0].text = str(index)
        row_cells[1].text = nfd.employee.name
        row_cells[2].text = nfd.employee.hr_number
        row_cells[3].text = nfd.employee.designation

        row_cells[4].text = (
            nfd.duty_date.strftime(
                "%d %b %Y"
            )
            if nfd.duty_date
            else "—"
        )

        row_cells[5].text = (
            nfd.duty_type
            or "—"
        )

    if not nfds:

        row_cells = nfd_table.add_row().cells

        row_cells[0].text = ""
        row_cells[1].text = (
            f"No NFD employees for {next_month_label}."
        )
        row_cells[2].text = ""
        row_cells[3].text = ""
        row_cells[4].text = ""
        row_cells[5].text = ""

    # ========================================================
    # Response
    # ========================================================

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )

    response["Content-Disposition"] = (
        f'attachment; '
        f'filename="Duty_Report_{selected_month}.docx"'
    )

    document.save(response)

    return response